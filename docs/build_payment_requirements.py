from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent
BLUE = '2E74B5'
DARK = '1F4D78'
INK = '0B2545'
MUTED = '5A6872'
LIGHT = 'F2F4F7'
CALLOUT = 'F4F6F9'
WHITE = 'FFFFFF'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    indent = OxmlElement('w:tblInd')
    indent.set(qn('w:w'), '120')
    indent.set(qn('w:type'), 'dxa')
    tbl_pr.append(indent)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), instruction)
    run._r.addnext(fld)


def base_doc(short_name, title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style, size, color, before, after in [
        ('Heading 1', 16, BLUE, 16, 8), ('Heading 2', 13, BLUE, 12, 6), ('Heading 3', 12, DARK, 8, 4)
    ]:
        s = doc.styles[style]
        s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color)
        s._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run(f'GEM & CRYSTAL FASHION HUB  |  {short_name}')
    set_font(r, 8, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Confidential — Internal planning document  |  Page ')
    set_font(r, 8, MUTED)
    add_field(footer, 'PAGE')

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(short_name)
    set_font(r, 10, DARK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, 24, INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(15)
    r = p.add_run(subtitle)
    set_font(r, 13, MUTED)

    meta = doc.add_table(rows=4, cols=2)
    set_table_widths(meta, [1800, 7560])
    for i, (label, value) in enumerate([
        ('Product', 'Gem & Crystal Fashion Hub'),
        ('Scope', 'Storefront, API, admin dashboard, POS, inventory and payments'),
        ('Version', '1.0 — pre-launch baseline'),
        ('Date', '4 September 2026'),
    ]):
        shade(meta.cell(i, 0), LIGHT)
        a = meta.cell(i, 0).paragraphs[0]; a.paragraph_format.space_after = Pt(0)
        b = meta.cell(i, 1).paragraphs[0]; b.paragraph_format.space_after = Pt(0)
        set_font(a.add_run(label), 10, INK, bold=True); set_font(b.add_run(value), 10, '202124')
    return doc


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), 11, '202124', bold=True)
        set_font(p.add_run(text[len(bold_prefix):]), 11, '202124')
    else:
        set_font(p.add_run(text), 11, '202124')
    return p


def bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.25)
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.167
        set_font(p.add_run(text), 11, '202124')


def numbered(doc, items):
    for text in items:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.25)
        p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.167
        set_font(p.add_run(text), 11, '202124')


def callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    shade(table.cell(0, 0), CALLOUT)
    p = table.cell(0, 0).paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + '  '), 10, DARK, bold=True)
    set_font(p.add_run(text), 10, '202124')


def matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), 9, INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(text), 9, '202124')
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def create_prd():
    doc = base_doc('PRODUCT REQUIREMENTS DOCUMENT', 'Gem & Crystal Fashion Hub Platform', 'Product requirements for the storefront, API, admin, POS and verified payments')
    heading(doc, '1. Product summary')
    para(doc, 'Gem & Crystal Fashion Hub is a unified fashion-commerce platform for a customer website and a physical shop. It includes catalogue browsing, variant-level inventory, online checkout, owner administration, cashier POS, and a shared API with PostgreSQL as the system of record.')
    callout(doc, 'Product principle:', 'One source of truth for products, stock, orders and payments. No customer screen, cashier action, or manually typed receipt can mark an order as paid.')
    heading(doc, '2. Problem and opportunity')
    para(doc, 'The shop needs a smooth customer journey and a trusted operational view. Without shared stock and payments, staff can oversell items, miss website payments, accept an unverified message, or waste time reconciling separate systems. The platform creates one operational flow while keeping customer checkout straightforward.')
    heading(doc, '3. Platform scope')
    matrix(doc, ['Area', 'Customer or staff outcome', 'Primary application'], [
        ('Storefront', 'Browse products, select variants, apply coupons and place an order.', 'Customer website'),
        ('Catalogue & inventory', 'Accurate variant prices and stock across all channels.', 'API + Admin'),
        ('Admin', 'Owner manages products, coupons, orders, stock and POS access.', 'Admin dashboard'),
        ('POS', 'Cashier scans or selects items, takes payment and completes in-store sales.', 'Dedicated POS'),
        ('Payments', 'M-Pesa payments are confirmed by the backend and visible in POS.', 'API + POS'),
    ], [1800, 4860, 2700])
    heading(doc, '4. Goals')
    bullets(doc, [
        'Send an M-Pesa payment prompt for eligible website orders and in-store sales.',
        'Show an automatic POS popup only after backend verification succeeds.',
        'Display the payment reference, amount, customer and M-Pesa receipt clearly.',
        'Prevent duplicate payment popups when M-Pesa retries a callback or the POS reconnects.',
        'Keep stock and payment status correct when a payment is cancelled or fails.',
        'Give the owner a usable admin workspace for product, inventory, coupon, order and POS control.',
        'Use the same database records for web orders and physical-store stock movements.',
    ])
    heading(doc, '5. Non-goals')
    bullets(doc, [
        'Accepting card payments or storing card details.',
        'Treating a customer-entered M-Pesa receipt as proof of payment.',
        'Replacing the Safaricom Daraja account, merchant onboarding, or settlement process.',
        'Building a customer-facing order-tracking portal beyond the existing order confirmation flow.',
    ])
    heading(doc, '6. Users and key journeys')
    matrix(doc, ['User', 'Need', 'Successful outcome'], [
        ('Website customer', 'Pay for an online order without sharing a receipt manually.', 'Receives M-Pesa prompt; order stays pending until confirmation.'),
        ('Counter customer', 'Pay at the shop from their phone.', 'Cashier sends prompt; sale completes only after verified result.'),
        ('Cashier', 'Know instantly that payment is safe to fulfil.', 'POS popup confirms the verified receipt and amount.'),
        ('Owner', 'See accurate paid and pending transactions.', 'Admin records reflect backend-confirmed payment state.'),
    ], [1550, 3250, 4560])
    heading(doc, '7. Functional requirements')
    matrix(doc, ['ID', 'Requirement', 'Priority'], [
        ('FR-00', 'Display active catalogue products with images, categories, variants, price and stock availability.', 'Must'),
        ('FR-00A', 'Maintain stock at variant level and use atomic updates for web orders and POS sales.', 'Must'),
        ('FR-00B', 'Allow owner-only administration of products, inventory, coupons, orders, settings and POS approvals.', 'Must'),
        ('FR-00C', 'Require cashier POS sessions to be approved by the owner and expire automatically.', 'Must'),
        ('FR-01', 'For an M-Pesa checkout, collect and validate the payer phone number.', 'Must'),
        ('FR-02', 'Create a pending order or POS sale before requesting M-Pesa payment.', 'Must'),
        ('FR-03', 'Send an STK payment request only from the backend after server-calculated totals are stored.', 'Must'),
        ('FR-04', 'Match the provider callback to the original checkout request and merchant request IDs.', 'Must'),
        ('FR-05', 'Verify amount, payer phone and M-Pesa receipt before setting payment status to PAID.', 'Must'),
        ('FR-06', 'Create one durable POS notification per confirmed M-Pesa payment.', 'Must'),
        ('FR-07', 'Poll the notification queue from an authenticated active POS session and show a popup.', 'Must'),
        ('FR-08', 'Acknowledge the popup without deleting the underlying payment record.', 'Should'),
        ('FR-09', 'On a confirmed failed/cancelled callback, mark payment FAILED and restore reserved stock.', 'Must'),
        ('FR-10', 'Prevent the same receipt or provider callback from creating multiple paid records.', 'Must'),
    ], [900, 6960, 1500])
    heading(doc, '8. Acceptance criteria')
    numbered(doc, [
        'A website M-Pesa order remains PENDING until the backend receives a valid successful callback.',
        'A physical POS M-Pesa sale prompts the customer and does not request that the cashier type a receipt code.',
        'After a valid callback, the POS displays one popup containing the reference, amount and M-Pesa receipt.',
        'A callback with the wrong amount, phone, request ID or secret does not mark payment as paid.',
        'A callback retry does not create another notification or change stock twice.',
        'A failed/cancelled callback restores the stock reserved for that transaction exactly once.',
        'If the POS is briefly offline, it receives the unacknowledged notification after reconnecting.',
        'A product or variant deactivated by the owner cannot be added to a new web order or POS sale.',
        'Stock is never reduced below zero by two simultaneous checkout attempts.',
        'Only owner accounts can access administration endpoints and approve POS sessions.',
    ])
    heading(doc, '9. Experience requirements')
    bullets(doc, [
        'The website says that payment is confirmed only after M-Pesa responds to the backend.',
        'The POS call-to-action is “Send M-Pesa Prompt,” not “Complete Sale,” until payment is confirmed.',
        'The popup is readable from a counter distance and requires an explicit cashier acknowledgement.',
        'If configuration is absent, the system explains that M-Pesa is not configured; it never pretends payment succeeded.',
    ])
    heading(doc, '10. Success measures')
    matrix(doc, ['Measure', 'Target at launch', 'How measured'], [
        ('Verified-payment popup delivery', '≥ 99% within 10 seconds of callback', 'API and POS event timestamps'),
        ('Duplicate payment alerts', '0', 'Unique payment notification records'),
        ('False paid status', '0', 'Audit review of callback validation failures'),
        ('Unresolved pending payments', '< 1% after daily review', 'Pending payment age report'),
        ('Inventory oversells', '0', 'Inventory movement and order reconciliation'),
    ], [3100, 2500, 3760])
    heading(doc, '11. Launch dependencies and risks')
    bullets(doc, [
        'Safaricom Daraja credentials, merchant shortcode/passkey, and an approved production setup are required.',
        'The callback address must be a public HTTPS endpoint reachable by Safaricom.',
        'Secrets must be held in the deployment secret manager, not in source control.',
        'Store staff need a short procedure for pending payments and an owner escalation path.',
    ])
    return doc


def create_trd():
    doc = base_doc('TECHNICAL REQUIREMENTS DOCUMENT', 'Gem & Crystal Platform Architecture', 'Technical requirements for storefront, API, administration, POS, inventory and payments')
    heading(doc, '1. Architecture overview')
    para(doc, 'The system consists of independent React applications for the storefront, owner admin dashboard and physical-store POS, backed by one Express API and PostgreSQL database. The API is the authority for authentication, product data, inventory, orders, payments and POS sessions. Safaricom Daraja returns final M-Pesa results to a protected API callback; the API validates the result inside a database transaction, then writes a durable notification for the POS.')
    matrix(doc, ['Component', 'Responsibility', 'Trust level'], [
        ('Storefront', 'Collect delivery and payer details; request checkout.', 'Untrusted client'),
        ('Admin', 'Owner-only product, inventory, coupon, order and POS control.', 'Authenticated owner client'),
        ('Dedicated POS / admin POS', 'Start counter sale; poll for confirmed payment alert.', 'Authenticated client'),
        ('API', 'Price calculation, STK request, callback validation, status changes.', 'Payment authority'),
        ('PostgreSQL', 'Orders, POS sales, inventory movements, notification queue, audit log.', 'System of record'),
        ('Safaricom Daraja', 'M-Pesa prompt and callback result.', 'External provider'),
    ], [1800, 5160, 2400])
    heading(doc, '2. Core platform flows')
    numbered(doc, [
        'Client submits product variant IDs, quantities, customer details and payer phone to the API.',
        'API validates input, calculates authoritative prices/totals, reserves stock and creates PENDING record.',
        'API requests an STK push from Daraja and stores CheckoutRequestID and MerchantRequestID.',
        'Customer approves or cancels payment on their phone.',
        'Daraja calls the public callback URL with the payment result.',
        'API validates callback secret, request IDs, result code, amount, payer phone and receipt.',
        'API updates payment status and creates one PaymentNotification record in the same transaction.',
        'Authenticated POS polls notification endpoint every three seconds, receives alert once and shows popup.',
    ])
    heading(doc, '3. Catalogue, stock and access control')
    bullets(doc, [
        'Product variants are the sellable unit. Each variant has an SKU, size, colour, price, optional sale price and stock quantity.',
        'API calculates checkout and POS line prices from current database variants; clients send only IDs and quantities.',
        'Stock updates use database conditions so stock cannot fall below zero during simultaneous sales.',
        'InventoryMovement records every sale, failed-payment return and owner adjustment with an actor and reference.',
        'Owner API routes require an owner JWT. Cashier POS sessions require a separate approved session token with an expiry.',
    ])
    heading(doc, '4. Data model requirements')
    matrix(doc, ['Entity', 'Required fields / constraints', 'Purpose'], [
        ('Order', 'paymentStatus, mpesaReceipt UNIQUE, checkout/merchant request IDs UNIQUE, initiated timestamp', 'Online payment lifecycle'),
        ('PosSale', 'paymentStatus, mpesaReceipt UNIQUE, checkout/merchant request IDs UNIQUE, initiated timestamp', 'Counter payment lifecycle'),
        ('PaymentNotification', 'orderId or posSaleId UNIQUE; paymentReference UNIQUE; acknowledgedAt', 'Durable one-time POS alert'),
        ('InventoryMovement', 'variant, quantity, previous/new stock, reference, actor', 'Audit stock reserve and release'),
        ('AuditLog', 'action, details, timestamp, source IP', 'Operational and security audit'),
    ], [1800, 4800, 2760])
    callout(doc, 'Transaction rule:', 'Payment status, stock release and notification creation must be committed atomically. A callback retry must be safe at every point.')
    heading(doc, '5. API requirements')
    matrix(doc, ['Endpoint', 'Authentication', 'Behaviour'], [
        ('POST /api/orders', 'Public with rate limit', 'Creates pending online order; starts STK only when configured.'),
        ('POST /api/pos/checkout', 'Active POS session bearer token', 'Calculates counter total server-side; starts STK for M-Pesa.'),
        ('POST /api/orders/mpesa-callback', 'Callback secret query token + strict payload validation', 'Processes verified M-Pesa result; never trusts client data.'),
        ('GET /api/pos/payment-notifications', 'Active POS session bearer token', 'Returns and acknowledges unacknowledged payment alerts.'),
        ('GET /api/orders/:orderNumber', 'One-time tracking token', 'Returns customer order status without exposing other orders.'),
        ('/api/products, /api/coupons, /api/settings', 'Public read / owner write as applicable', 'Catalogue, promotion and store configuration services.'),
        ('/api/admin and protected owner routes', 'Owner bearer token', 'Administration, reporting, inventory and POS authorisation.'),
    ], [3000, 2600, 3760])
    heading(doc, '6. Callback validation rules')
    bullets(doc, [
        'Reject callback unless M-Pesa integration is configured and the callback secret matches in constant time.',
        'Accept only a strict, expected callback shape; reject malformed bodies.',
        'Look up transaction by CheckoutRequestID and require matching MerchantRequestID.',
        'For successful callbacks, require a valid receipt, exact order amount and normalized payer phone match.',
        'Treat mismatches as review cases; do not mark paid or emit a POS alert.',
        'Enforce unique receipt/request IDs so duplicate deliveries cannot pay twice.',
    ])
    heading(doc, '7. Failure and recovery behaviour')
    matrix(doc, ['Situation', 'System response', 'Owner action'], [
        ('Customer cancels / provider reports failure', 'Mark FAILED; restore reserved stock once; write audit log.', 'Cashier may retry only with a new payment request.'),
        ('Callback data mismatch', 'Leave payment PENDING; write review audit event; no alert.', 'Owner verifies against provider before manual resolution.'),
        ('POS offline after confirmation', 'Keep notification unacknowledged in database.', 'POS receives popup when it reconnects.'),
        ('Daraja unavailable before request starts', 'Keep pending record; do not claim payment success.', 'Owner reviews or retries using controlled process.'),
        ('Duplicate callback', 'Detect PAID/unique receipt and return idempotent result.', 'No action required.'),
    ], [2400, 4400, 2560])
    heading(doc, '8. Security and configuration')
    bullets(doc, [
        'Use HTTPS for production callback URL and restrict API CORS origins to deployed storefront/admin/POS domains.',
        'Store MPESA_CONSUMER_KEY, MPESA_CONSUMER_SECRET, MPESA_PASSKEY and MPESA_CALLBACK_SECRET only in a secret manager.',
        'Use independent strong secrets for JWT, POS sessions and M-Pesa callback verification.',
        'Do not log credentials, callback secrets or full access tokens.',
        'Rate-limit public checkout routes; cap payload sizes and validate every input with strict schemas.',
        'Keep browser admin/POS tokens in memory and protect POS notification polling with an active session token.',
    ])
    heading(doc, '9. Required environment variables')
    matrix(doc, ['Variable', 'Required in production', 'Purpose'], [
        ('MPESA_ENV', 'Yes', 'sandbox or production provider endpoint selection'),
        ('MPESA_CONSUMER_KEY / MPESA_CONSUMER_SECRET', 'Yes', 'Daraja OAuth credentials'),
        ('MPESA_SHORTCODE / MPESA_PASSKEY', 'Yes', 'STK request identity and password'),
        ('MPESA_CALLBACK_URL', 'Yes', 'Public HTTPS callback endpoint'),
        ('MPESA_CALLBACK_SECRET', 'Yes', 'Protect callback URL from unauthorised requests'),
        ('DATABASE_URL, JWT_SECRET, POS_SESSION_SECRET', 'Yes', 'Data access and application sessions'),
    ], [3400, 2200, 3760])
    heading(doc, '10. Observability and operations')
    bullets(doc, [
        'Log structured events for payment initiated, confirmed, failed, mismatch/review and notification delivered.',
        'Monitor pending payments by age and alert owner if a payment remains pending beyond the agreed window.',
        'Record audit entries for provider callback outcomes and manual owner status changes.',
        'Back up PostgreSQL and protect migration/deployment history before launch.',
    ])
    heading(doc, '11. Test plan')
    matrix(doc, ['Test', 'Expected result'], [
        ('Successful website payment callback', 'Order PAID, receipt stored, one POS popup notification.'),
        ('Successful physical POS payment callback', 'Sale PAID, receipt stored, one POS popup notification.'),
        ('Wrong request ID, amount or phone', 'Payment remains PENDING; audit entry created; no popup.'),
        ('Cancelled callback', 'Payment FAILED and stock restored once.'),
        ('Duplicate callback', 'No duplicate receipt, status update, stock movement or popup.'),
        ('POS reconnect', 'Unacknowledged notification appears after authenticated reconnect.'),
        ('No M-Pesa configuration', 'API reports unavailable; never reports payment as paid.'),
        ('Concurrent web/POS stock request', 'One request fails cleanly when remaining stock is insufficient.'),
        ('Unauthorised admin request', 'API returns 401/403 without exposing protected records.'),
    ], [4000, 5360])
    heading(doc, '12. Deployment checklist')
    numbered(doc, [
        'Create and test Daraja sandbox application; register the production callback URL during go-live.',
        'Set production secrets through deployment configuration and confirm no secrets are committed.',
        'Apply Prisma schema changes and take a database backup.',
        'Deploy API before storefront/admin/POS builds.',
        'Perform sandbox payment, cancellation, callback retry and reconnect tests.',
        'Train cashiers: wait for the verified POS popup before handing over goods.',
        'Review payment and pending-order dashboard daily during launch week.',
    ])
    return doc


if __name__ == '__main__':
    create_prd().save(OUT / 'Gem-Crystal-Platform-PRD.docx')
    create_trd().save(OUT / 'Gem-Crystal-Platform-TRD.docx')
